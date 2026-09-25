#!/usr/bin/env python3
"""Import Google Forms CSV export into analysis workbook raw responses sheet."""

import csv
import re
import sys
from pathlib import Path

import openpyxl

SURVEYS_DIR = Path(__file__).resolve().parent
ANALYSIS_FILE = SURVEYS_DIR / "DUP_analiz_otvetov_oprosnika.xlsx"
RAW_SHEET = "Сырые ответы"

# Map Google Forms question text prefixes to question IDs
ID_PATTERNS = [
    (r"регион|филиал", "A1"),
    (r"должность", "A2"),
    (r"действующих проектов", "A3"),
    (r"сотрудников в проектной", "A4"),
    (r"взаимодействуете с куратором", "A5"),
    (r"закреплённым куратором", "A6"),
    (r"понятен вам процесс запуска нового проекта", "B1"),
    (r"получаете новый проект", "B2"),
    (r"полным приходит пакет документов", "B3"),
    (r"документы чаще всего отсутствуют", "B4"),
    (r"подготовка проектной документации", "B5"),
    (r"согласовать бюджет проекта с Финблоком", "B6"),
    (r"закупочные процедуры.*задерживают", "B7"),
    (r"длительность закупочной процедуры", "B8"),
    (r"статус-встречи по проектам реализации", "B9"),
    (r"инициирует и ведёт статус-встречи", "B10"),
    (r"блокера на этапе реализации", "B11"),
    (r"сопровождения действующих проектов", "C1"),
    (r"сопровождения выполняет филиал самостоятельно", "C2"),
    (r"ДУП реагирует на операционные проблемы", "C3"),
    (r"до полной готовности.*пилот", "C4"),
    (r"к чему это приводило", "C5"),
    (r"дефектов компонентов", "C6"),
    (r"продление или модернизацию контракта", "C7"),
    (r"боль в операционном сопровождении", "C8"),
    (r"формирует реестр платежей", "D1"),
    (r"приоритизации платежей", "D2"),
    (r"платежи по вашим заявкам задерживаются", "D3"),
    (r"причина задержки оплат", "D4"),
    (r"прогнозе поступлений доходов", "D5"),
    (r"прогнозу поступлений/расходов", "D6"),
    (r"нехватки бюджета проекта", "D7"),
    (r"как решался вопрос", "D8"),
    (r"разграничены зоны ответственности", "E1"),
    (r"к кому в ДУП обращаться", "E2"),
    (r"качество коммуникации с куратором", "E3"),
    (r"каналы связи с ДУП", "E4"),
    (r"эскалации проблем в обход ДУП", "E5"),
    (r"открытии нового филиала", "E6"),
    (r"понятен был процесс и кто координировал", "E7"),
    (r"улучшили во взаимодействии с ДУП", "E8"),
    (r"взаимодействие с ДРБ", "F1_ДРБ"),
    (r"взаимодействие с Финблоком", "F1_Финблок"),
    (r"взаимодействие с Юр блоком", "F1_Юр_блок"),
    (r"Закупками \(ОЛЖА\)", "F1_Закупки_(ОЛЖ"),
    (r"СТ \(тех", "F1_СТ_(тех._блок"),
    (r"Зерек", "F1_Зерек_/_Zerek"),
    (r"взаимодействие с HR", "F1_HR"),
    (r"Лабораторией", "F1_Лаборатория_(п"),
    (r"задержки и недопонимание", "F2"),
    (r"ДРБ передаёт проекты", "F3"),
    (r"Юр блок сопровождает", "F4"),
    (r"проблему во взаимодействии с центральным офисом", "F5"),
    (r"инструменты вы используете", "G1"),
    (r"платформа портфеля проектов", "G2"),
    (r"не хватает в текущих IT", "G3"),
    (r"ручной сбор данных и отчётность", "G4"),
    (r"ТОП-3 проблемы", "H1"),
    (r"ТОП-3 процесса.*автоматизировать", "H2"),
    (r"изменить одну вещь", "H3"),
    (r"углублённом интервью", "H4"),
    (r"Контакт для связи", "H5"),
]


def match_id(header: str) -> str | None:
    h = header.lower().strip()
    for pattern, qid in ID_PATTERNS:
        if re.search(pattern, header, re.I):
            return qid
    return None


def load_header_map(ws):
    mapping = {}
    for col in range(5, ws.max_column + 1):
        val = ws.cell(row=1, column=col).value
        if not val:
            continue
        qid = val.split(":", 1)[0].strip()
        mapping[qid] = col
    return mapping


def import_csv(csv_path: Path):
    if not ANALYSIS_FILE.exists():
        raise SystemExit(f"Analysis workbook not found: {ANALYSIS_FILE}. Run build_analysis_workbook.py first.")

    wb = openpyxl.load_workbook(ANALYSIS_FILE)
    ws = wb[RAW_SHEET]
    col_map = load_header_map(ws)

    # Map CSV headers to workbook columns via question IDs
    csv_to_col = {}
    with csv_path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames:
            raise SystemExit("CSV has no headers")
        for header in reader.fieldnames:
            if header.lower().startswith("отметка времени") or header.lower() == "timestamp":
                csv_to_col[header] = 1  # date column
                continue
            if "email" in header.lower():
                continue
            qid = match_id(header)
            if qid and qid in col_map:
                csv_to_col[header] = col_map[qid]
            else:
                print(f"Warning: unmapped column: {header[:60]}...")

        start_row = 2
        while ws.cell(row=start_row, column=2).value:
            start_row += 1

        count = 0
        for row_data in reader:
            row_idx = start_row + count
            for header, value in row_data.items():
                col = csv_to_col.get(header)
                if col:
                    ws.cell(row=row_idx, column=col, value=value.strip() if value else "")
            count += 1

    wb.save(ANALYSIS_FILE)
    print(f"Imported {count} responses into {ANALYSIS_FILE} (sheet '{RAW_SHEET}', from row {start_row})")


def main():
    if len(sys.argv) < 2:
        print("Usage: python3 import_survey_responses.py <google_forms_export.csv>")
        sys.exit(1)
    import_csv(Path(sys.argv[1]))


if __name__ == "__main__":
    main()
