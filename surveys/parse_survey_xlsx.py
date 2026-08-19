#!/usr/bin/env python3
"""Parse corrected DUP survey from Excel sheet 'Анкета для заполнения'."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

import openpyxl

QuestionType = Literal["text", "paragraph", "radio", "checkbox", "scale", "grid"]

OPTION_RE = re.compile(r"^\s*☐\s*(.+)$")
SCALE_RE = re.compile(r"^\s*Ответ:\s*1\s+2\s+3\s+4\s+5\s*$")
TEXT_RE = re.compile(r"^\s*Ответ:\s*_{3,}")
GRID_ROW_RE = re.compile(r"^\s*•\s*(.+?):\s*1\s+2\s+3\s+4\s+5\s*$")
QUESTION_ID_RE = re.compile(r"^([A-H]\d+)\.\s*$")
SECTION_RE = re.compile(r"^Блок\s+[A-H]\.\s*.+$")

DEFAULT_XLSX = Path(__file__).resolve().parent / "Опросник_директоров_Филиалов.xlsx"
SHEET_NAME = "Анкета для заполнения"

OPTIONAL_IDS = {"C5", "D8", "H5"}


@dataclass
class Question:
    qid: str
    text: str
    qtype: QuestionType
    options: list[str] = field(default_factory=list)
    required: bool = True
    section: str = ""


@dataclass
class Section:
    title: str
    questions: list[Question] = field(default_factory=list)


@dataclass
class Survey:
    title: str
    description: str
    sections: list[Section] = field(default_factory=list)


def _is_paragraph(text: str) -> bool:
    markers = ("Опишите", "Назовите", "Если да", "ТОП-3", "изменить одну вещь")
    return any(m in text for m in markers)


def parse_survey_xlsx(path: Path | str = DEFAULT_XLSX) -> Survey:
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Лист «{SHEET_NAME}» не найден в {path}")
    ws = wb[SHEET_NAME]

    survey = Survey(
        title="ДУП — опрос директоров филиалов",
        description=(
            "Диагностика процессов Департамента управления проектами (ДУП), ТОО «Көркем Телеком». "
            "Цель — выявить боли во взаимодействии филиал ↔ ДУП ↔ центральный офис. "
            "Время заполнения: 20–30 минут."
        ),
    )

    current_section: Section | None = None
    current_q: Question | None = None

    def flush_question() -> None:
        nonlocal current_q
        if current_q and current_section:
            if current_q.qid == "G3":
                current_q.qtype = "checkbox"
            current_section.questions.append(current_q)
        current_q = None

    for row in ws.iter_rows(min_row=1, max_col=2, values_only=True):
        col_a, col_b = (row + (None, None))[:2]
        a = str(col_a).strip() if col_a is not None else ""
        b = str(col_b).strip() if col_b is not None else ""

        if not a and not b:
            if current_q and current_q.options and current_q.qtype in ("radio", "checkbox", "grid"):
                flush_question()
            continue

        if a and SECTION_RE.match(a):
            flush_question()
            current_section = Section(title=a)
            survey.sections.append(current_section)
            continue

        if not current_section:
            continue

        if a and QUESTION_ID_RE.match(a):
            flush_question()
            qid = QUESTION_ID_RE.match(a).group(1)
            current_q = Question(
                qid=qid,
                text=b,
                qtype="radio",
                required=qid not in OPTIONAL_IDS,
                section=current_section.title,
            )
            continue

        if not current_q or not b:
            continue

        if m := OPTION_RE.match(b):
            current_q.options.append(m.group(1).strip())
            current_q.qtype = "checkbox" if "(можно несколько)" in current_q.text else "radio"
            continue

        if SCALE_RE.match(b):
            current_q.qtype = "scale"
            flush_question()
            continue

        if TEXT_RE.match(b):
            current_q.qtype = "paragraph" if _is_paragraph(current_q.text) else "text"
            flush_question()
            continue

        if current_q.qid == "F1" and (m := GRID_ROW_RE.match(b)):
            current_q.qtype = "grid"
            current_q.options.append(m.group(1).strip())

    flush_question()

    return survey


def question_count(survey: Survey) -> int:
    return sum(len(s.questions) for s in survey.sections)


if __name__ == "__main__":
    s = parse_survey_xlsx()
    print(f"{s.title}: {len(s.sections)} блоков, {question_count(s)} вопросов")
    for sec in s.sections:
        print(f"\n{sec.title}")
        for q in sec.questions:
            opts = f" [{len(q.options)} opt]" if q.options else ""
            print(f"  {q.qid} ({q.qtype}) {q.text[:70]}{opts}")
