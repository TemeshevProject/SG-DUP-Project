#!/usr/bin/env python3
"""Add demo response rows to analysis workbook for testing dashboards."""

import random
from datetime import date, timedelta
from pathlib import Path

import openpyxl

ANALYSIS = Path(__file__).resolve().parent / "DUP_analiz_otvetov_oprosnika.xlsx"
RAW = "Сырые ответы"

DEMO_FILIALS = [
    ("Алматы", "Директор филиала"),
    ("Туркестан", "Главный инженер"),
    ("Астана", "Региональный менеджер проекта"),
    ("Тараз", "Директор филиала"),
    ("Шымкент", "Заместитель директора филиала"),
]

SCALE_COLS_PREFIX = [
    "B1", "B3", "B6", "C1", "C3", "C6", "D2", "D6", "E1", "E2", "E3", "F3", "F4", "G2",
]


def col_map(ws):
    m = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(1, c).value
        if h and ":" in str(h):
            m[str(h).split(":", 1)[0].strip()] = c
    return m


def main():
    wb = openpyxl.load_workbook(ANALYSIS)
    ws = wb[RAW]
    cmap = col_map(ws)
    start = 2
    while ws.cell(start, 2).value and not str(ws.cell(start, 2).value).startswith("("):
        start += 1

    for i, (filial, role) in enumerate(DEMO_FILIALS):
        r = start + i
        ws.cell(r, 1, value=(date.today() - timedelta(days=i * 2)).isoformat())
        ws.cell(r, 2, value=filial)
        ws.cell(r, 3, value=role)
        for qid in SCALE_COLS_PREFIX:
            if qid in cmap:
                ws.cell(r, cmap[qid], value=random.randint(2, 4))
        for qid in [k for k in cmap if k.startswith("F1_")]:
            ws.cell(r, cmap[qid], value=random.randint(2, 5))
        if "H1" in cmap:
            ws.cell(r, cmap["H1"], value="Задержки закупок; неполный пакет от ДРБ; хаос в прогнозе поступлений")
        if "H4" in cmap:
            ws.cell(r, cmap["H4"], value=random.choice(["Да", "Возможно, по согласованию", "Нет"]))

    wb.save(ANALYSIS)
    print(f"Added {len(DEMO_FILIALS)} demo rows to {ANALYSIS} (rows {start}-{start+len(DEMO_FILIALS)-1})")
    print("Удалите демо-строки после импорта реальных ответов.")


if __name__ == "__main__":
    main()
