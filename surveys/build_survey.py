#!/usr/bin/env python3
"""Generate DUP branch director survey Excel workbook."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTPUT = "/workspace/surveys/DUP_oprosnik_direktorov_filialov.xlsx"

# (block_id, block_title, questions)
# question: dict with keys: id, text, type, options (list or None), required
BLOCKS = [
    (
        "A",
        "Блок A. Общая информация о филиале",
        [
            {
                "id": "A1",
                "text": "Ваш регион / филиал",
                "type": "текст",
                "options": None,
                "required": True,
            },
            {
                "id": "A2",
                "text": "Ваша должность",
                "type": "один из списка",
                "options": [
                    "Директор филиала",
                    "Заместитель директора филиала",
                    "Главный инженер",
                    "Региональный менеджер проекта",
                    "Менеджер проектов филиала",
                    "Другое (укажите в комментарии)",
                ],
                "required": True,
            },
            {
                "id": "A3",
                "text": "Сколько действующих проектов ГЧП/ГЗ сейчас ведёт филиал?",
                "type": "один из списка",
                "options": ["1–2", "3–5", "6–10", "более 10"],
                "required": True,
            },
            {
                "id": "A4",
                "text": "Сколько сотрудников в проектной/операционной команде филиала (приблизительно)?",
                "type": "один из списка",
                "options": ["до 5", "6–15", "16–30", "более 30"],
                "required": True,
            },
            {
                "id": "A5",
                "text": "Как часто вы взаимодействуете с куратором/менеджером проекта со стороны ДУП?",
                "type": "один из списка",
                "options": ["Ежедневно", "Несколько раз в неделю", "Раз в неделю", "Реже раза в неделю", "По мере необходимости"],
                "required": True,
            },
            {
                "id": "A6",
                "text": "Знаете ли вы, кто является вашим закреплённым куратором в ДУП (ФИО)?",
                "type": "да/нет",
                "options": ["Да", "Нет", "Знаю частично / куратор менялся"],
                "required": True,
            },
        ],
    ),
    (
        "B",
        "Блок B. Реализация нового проекта",
        [
            {
                "id": "B1",
                "text": "Насколько понятен вам процесс запуска нового проекта в филиале (от получения поручения до сдачи объекта)?",
                "type": "шкала 1–5",
                "options": ["1 — совсем непонятен", "2", "3", "4", "5 — полностью понятен и формализован"],
                "required": True,
            },
            {
                "id": "B2",
                "text": "Как вы обычно получаете новый проект от ДУП / ДРБ?",
                "type": "несколько из списка",
                "options": [
                    "Карточка проекта в SimBase / корпоративной системе",
                    "Служебная записка / приказ",
                    "Сообщение в мессенджере (WhatsApp, Telegram и т.д.)",
                    "Звонок / видеосвязь",
                    "Письмо на email",
                    "Другое",
                ],
                "required": True,
            },
            {
                "id": "B3",
                "text": "Достаточно ли полным приходит пакет документов при старте проекта (договор, ФЭМ, ТЗ, адресная программа и т.д.)?",
                "type": "шкала 1–5",
                "options": ["1 — почти всегда неполный", "2", "3", "4", "5 — почти всегда полный"],
                "required": True,
            },
            {
                "id": "B4",
                "text": "Какие документы чаще всего отсутствуют или приходят с задержкой? (можно несколько)",
                "type": "несколько из списка",
                "options": [
                    "Адресная программа",
                    "Финансово-экономическая модель (ФЭМ)",
                    "Техническое задание / тех.карты",
                    "Утверждённый бюджет",
                    "Договор с госпартнёром",
                    "Приказ на реализацию",
                    "Ничего критичного не отсутствует",
                    "Другое",
                ],
                "required": False,
            },
            {
                "id": "B5",
                "text": "Сколько в среднем занимает подготовка проектной документации филиалом (устав, план, приказ)?",
                "type": "один из списка",
                "options": ["до 1 недели", "1–2 недели", "2–4 недели", "более месяца", "Сроки непредсказуемы"],
                "required": True,
            },
            {
                "id": "B6",
                "text": "Насколько сложно согласовать бюджет проекта с Финблоком через ДУП?",
                "type": "шкала 1–5",
                "options": ["1 — очень сложно, частые задержки", "2", "3", "4", "5 — проходит быстро и предсказуемо"],
                "required": True,
            },
            {
                "id": "B7",
                "text": "Как часто закупочные процедуры (выше упрощённого порога) задерживают реализацию проекта?",
                "type": "один из списка",
                "options": [
                    "Почти никогда",
                    "Иногда (менее 25% проектов)",
                    "Часто (25–50% проектов)",
                    "Очень часто (более 50% проектов)",
                    "Не могу оценить",
                ],
                "required": True,
            },
            {
                "id": "B8",
                "text": "Оцените среднюю длительность закупочной процедуры выше упрощённого порога (от заявки до договора)",
                "type": "один из списка",
                "options": ["до 2 недель", "2–4 недели", "1–2 месяца", "более 2 месяцев", "Сильно варьируется"],
                "required": True,
            },
            {
                "id": "B9",
                "text": "Проводятся ли регулярные статус-встречи по проектам реализации?",
                "type": "один из списка",
                "options": [
                    "Да, еженедельно",
                    "Да, раз в две недели",
                    "Да, ежемесячно",
                    "Нерегулярно",
                    "Практически не проводятся",
                ],
                "required": True,
            },
            {
                "id": "B10",
                "text": "Кто чаще всего инициирует и ведёт статус-встречи по реализации?",
                "type": "один из списка",
                "options": [
                    "Куратор ДУП",
                    "Менеджер проектов филиала",
                    "Директор филиала",
                    "Никто системно не ведёт",
                    "Другое",
                ],
                "required": True,
            },
            {
                "id": "B11",
                "text": "Опишите 1–2 типичных блокера на этапе реализации в вашем регионе",
                "type": "развёрнутый текст",
                "options": None,
                "required": False,
            },
        ],
    ),
    (
        "C",
        "Блок C. Постпроектное сопровождение",
        [
            {
                "id": "C1",
                "text": "Насколько понятен процесс сопровождения действующих проектов после ввода в эксплуатацию?",
                "type": "шкала 1–5",
                "options": ["1 — непонятен", "2", "3", "4", "5 — полностью понятен"],
                "required": True,
            },
            {
                "id": "C2",
                "text": "Какие задачи сопровождения выполняет филиал самостоятельно без участия ДУП? (можно несколько)",
                "type": "несколько из списка",
                "options": [
                    "Контроль работоспособности оборудования",
                    "Сопровождение канала связи",
                    "Взаимодействие с госпартнёром на местах",
                    "Организация поверки (координация с лабораторией)",
                    "Перенос оборудования по запросу госпартнёра",
                    "Подготовка ежемесячной отчётности госпартнёру",
                    "Инициация продления / модернизации контракта",
                    "Другое",
                ],
                "required": True,
            },
            {
                "id": "C3",
                "text": "Насколько своевременно ДУП реагирует на операционные проблемы (инциденты, простои, претензии госпартнёра)?",
                "type": "шкала 1–5",
                "options": ["1 — очень медленно", "2", "3", "4", "5 — оперативно"],
                "required": True,
            },
            {
                "id": "C4",
                "text": "Сталкивались ли вы с ситуацией, когда продукт/оборудование передали в эксплуатацию до полной готовности (пилот не завершён)?",
                "type": "да/нет",
                "options": ["Да, неоднократно", "Да, единичные случаи", "Нет", "Затрудняюсь ответить"],
                "required": True,
            },
            {
                "id": "C5",
                "text": "Если да — к чему это приводило? (можно несколько)",
                "type": "несколько из списка",
                "options": [
                    "Претензии госпартнёра",
                    "Штрафы / санкции",
                    "Дополнительные работы и перерасход бюджета",
                    "Репутационные потери",
                    "Не приводило к серьёзным последствиям",
                    "Не применимо",
                ],
                "required": False,
            },
            {
                "id": "C6",
                "text": "Насколько налажен процесс передачи информации о дефектах компонентов (от монтажников/эксплуатации к закупкам/ТЗ)?",
                "type": "шкала 1–5",
                "options": ["1 — не налажен вообще", "2", "3", "4", "5 — работает системно"],
                "required": True,
            },
            {
                "id": "C7",
                "text": "Как часто филиал инициирует продление или модернизацию контракта с госпартнёром?",
                "type": "один из списка",
                "options": [
                    "Регулярно, по регламенту",
                    "Иногда, когда срок подходит",
                    "Редко — ДУП/ДРБ инициируют сами",
                    "Процесс не формализован",
                ],
                "required": True,
            },
            {
                "id": "C8",
                "text": "Опишите главную боль в операционном сопровождении в вашем филиале",
                "type": "развёрнутый текст",
                "options": None,
                "required": False,
            },
        ],
    ),
    (
        "D",
        "Блок D. Финансы и платежи",
        [
            {
                "id": "D1",
                "text": "Кто в филиале формирует реестр платежей?",
                "type": "один из списка",
                "options": [
                    "Менеджер проектов / проектный офис филиала",
                    "Бухгалтерия филиала",
                    "Директор филиала",
                    "Несколько ролей совместно",
                    "Другое",
                ],
                "required": True,
            },
            {
                "id": "D2",
                "text": "Насколько понятен процесс приоритизации платежей (филиал → ДУП → Финблок)?",
                "type": "шкала 1–5",
                "options": ["1 — непонятен", "2", "3", "4", "5 — полностью понятен"],
                "required": True,
            },
            {
                "id": "D3",
                "text": "Как часто платежи по вашим заявкам задерживаются дольше критичного срока?",
                "type": "один из списка",
                "options": [
                    "Почти никогда",
                    "Иногда (до 25% заявок)",
                    "Часто (25–50%)",
                    "Очень часто (более 50%)",
                ],
                "required": True,
            },
            {
                "id": "D4",
                "text": "Какова типичная причина задержки оплат? (можно несколько)",
                "type": "несколько из списка",
                "options": [
                    "Неполный пакет документов от филиала",
                    "Долгое согласование в ДУП",
                    "Очередь / приоритизация в Финблоке",
                    "Проблемы с закупочной процедурой",
                    "Нехватка бюджета по проекту",
                    "Не знаю причину",
                    "Другое",
                ],
                "required": True,
            },
            {
                "id": "D5",
                "text": "Участвует ли филиал в прогнозе поступлений доходов от госпартнёра?",
                "type": "один из списка",
                "options": [
                    "Да, регулярно по шаблону",
                    "Да, но нерегулярно / по запросу в чате",
                    "Нет, это делает только ДУП",
                    "Процесс не налажен",
                ],
                "required": True,
            },
            {
                "id": "D6",
                "text": "Насколько удобен текущий формат обмена данными по прогнозу поступлений/расходов с Финблоком?",
                "type": "шкала 1–5",
                "options": ["1 — крайне неудобен (хаос в чатах)", "2", "3", "4", "5 — удобен и структурирован"],
                "required": True,
            },
            {
                "id": "D7",
                "text": "Сталкивались ли с ситуацией нехватки бюджета проекта (перерасход по ФЭМ)?",
                "type": "да/нет",
                "options": ["Да, неоднократно", "Да, единичные случаи", "Нет", "Не знаю"],
                "required": True,
            },
            {
                "id": "D8",
                "text": "Если да — как решался вопрос?",
                "type": "развёрнутый текст",
                "options": None,
                "required": False,
            },
        ],
    ),
    (
        "E",
        "Блок E. Взаимодействие с ДУП",
        [
            {
                "id": "E1",
                "text": "Насколько чётко разграничены зоны ответственности между филиалом и ДУП?",
                "type": "шкала 1–5",
                "options": ["1 — полностью размыты", "2", "3", "4", "5 — чётко зафиксированы и соблюдаются"],
                "required": True,
            },
            {
                "id": "E2",
                "text": "Понимаете ли вы, к кому в ДУП обращаться по разным типам вопросов (реализация / сопровождение / финансы / методология)?",
                "type": "шкала 1–5",
                "options": ["1 — не понимаю", "2", "3", "4", "5 — полностью понимаю"],
                "required": True,
            },
            {
                "id": "E3",
                "text": "Как вы оцениваете качество коммуникации с куратором ДУП?",
                "type": "шкала 1–5",
                "options": ["1 — очень плохое", "2", "3", "4", "5 — отличное"],
                "required": True,
            },
            {
                "id": "E4",
                "text": "Какие каналы связи с ДУП используются чаще всего? (можно несколько)",
                "type": "несколько из списка",
                "options": [
                    "Корпоративная почта",
                    "Мессенджеры (WhatsApp, Telegram)",
                    "SimBase / корпоративная система",
                    "Jira / задачи",
                    "Видеосвязь (Zoom, Teams)",
                    "Телефонные звонки",
                    "Другое",
                ],
                "required": True,
            },
            {
                "id": "E5",
                "text": "Бывали ли случаи эскалации проблем в обход ДУП напрямую к топ-менеджменту или внешним партнёрам?",
                "type": "да/нет",
                "options": ["Да, со стороны филиала", "Да, со стороны внешних партнёров/госорганов", "Да, с обеих сторон", "Нет"],
                "required": True,
            },
            {
                "id": "E6",
                "text": "Участвовал ли филиал в открытии нового филиала / расширении присутствия в регионе?",
                "type": "да/нет",
                "options": ["Да", "Нет", "В процессе"],
                "required": True,
            },
            {
                "id": "E7",
                "text": "Если да — насколько понятен был процесс и кто координировал со стороны компании?",
                "type": "развёрнутый текст",
                "options": None,
                "required": False,
            },
            {
                "id": "E8",
                "text": "Что бы вы улучшили во взаимодействии с ДУП в первую очередь?",
                "type": "развёрнутый текст",
                "options": None,
                "required": True,
            },
        ],
    ),
    (
        "F",
        "Блок F. Взаимодействие с другими подразделениями ЦО",
        [
            {
                "id": "F1",
                "text": "Оцените взаимодействие филиала с подразделениями (1 — очень плохо, 5 — отлично)",
                "type": "матрица",
                "options": ["ДРБ", "Финблок", "Юр блок", "Закупки (ОЛЖА)", "СТ (тех. блок)", "Зерек / Zerek Road tech", "HR", "Лаборатория (поверка)"],
                "required": True,
            },
            {
                "id": "F2",
                "text": "С каким подразделением ЦО чаще всего возникают задержки и недопонимание?",
                "type": "несколько из списка",
                "options": [
                    "ДРБ",
                    "Финблок",
                    "Юр блок",
                    "Закупки",
                    "СТ",
                    "Зерек",
                    "HR",
                    "Нет системных проблем",
                    "Другое",
                ],
                "required": True,
            },
            {
                "id": "F3",
                "text": "Насколько своевременно ДРБ передаёт проекты с полным пакетом документов?",
                "type": "шкала 1–5",
                "options": ["1 — почти никогда", "2", "3", "4", "5 — почти всегда своевременно"],
                "required": True,
            },
            {
                "id": "F4",
                "text": "Насколько оперативно Юр блок сопровождает договоры и открытие филиала?",
                "type": "шкала 1–5",
                "options": ["1 — очень медленно", "2", "3", "4", "5 — оперативно"],
                "required": True,
            },
            {
                "id": "F5",
                "text": "Опишите типичную проблему во взаимодействии с центральным офисом (кроме ДУП)",
                "type": "развёрнутый текст",
                "options": None,
                "required": False,
            },
        ],
    ),
    (
        "G",
        "Блок G. Инструменты и системы",
        [
            {
                "id": "G1",
                "text": "Какие инструменты вы используете в ежедневной проектной работе? (можно несколько)",
                "type": "несколько из списка",
                "options": [
                    "SimBase / корпоративная платформа проектов",
                    "Jira",
                    "Confluence",
                    "Excel / Google Таблицы",
                    "1С",
                    "Мессенджеры",
                    "Бумажный документооборот",
                    "Другое",
                ],
                "required": True,
            },
            {
                "id": "G2",
                "text": "Насколько полноценно работает платформа портфеля проектов (документация, бюджет план/факт, риски, задачи)?",
                "type": "шкала 1–5",
                "options": ["1 — не используем / не работает", "2", "3", "4", "5 — используем полноценно"],
                "required": True,
            },
            {
                "id": "G3",
                "text": "Чего не хватает в текущих IT-инструментах для вашей работы?",
                "type": "несколько из списка",
                "options": [
                    "Единая карточка проекта с обязательными полями",
                    "Автоуведомления (срок контракта, просрочки, платежи)",
                    "Интеграция закупок и 1С",
                    "Мониторинг оборудования (Zabbix и т.д.)",
                    "Шаблоны отчётов и реестров",
                    "Мобильный доступ",
                    "Ничего критичного",
                    "Другое",
                ],
                "required": True,
            },
            {
                "id": "G4",
                "text": "Сколько времени в неделю уходит на ручной сбор данных и отчётность (вне системы)?",
                "type": "один из списка",
                "options": ["менее 2 часов", "2–5 часов", "5–10 часов", "более 10 часов"],
                "required": True,
            },
        ],
    ),
    (
        "H",
        "Блок H. Итоговая оценка и предложения",
        [
            {
                "id": "H1",
                "text": "Назовите ТОП-3 проблемы в процессах, которые сильнее всего мешают работе филиала",
                "type": "развёрнутый текст",
                "options": None,
                "required": True,
            },
            {
                "id": "H2",
                "text": "Назовите ТОП-3 процесса, которые нужно автоматизировать в первую очередь",
                "type": "развёрнутый текст",
                "options": None,
                "required": True,
            },
            {
                "id": "H3",
                "text": "Если бы вы могли изменить одну вещь в работе ДУП и филиалов — что бы это было?",
                "type": "развёрнутый текст",
                "options": None,
                "required": True,
            },
            {
                "id": "H4",
                "text": "Готовы ли вы участвовать в 30–45-минутном углублённом интервью для детализации ответов?",
                "type": "да/нет",
                "options": ["Да", "Нет", "Возможно, по согласованию"],
                "required": True,
            },
            {
                "id": "H5",
                "text": "Контакт для связи (email / телефон) — если готовы к интервью",
                "type": "текст",
                "options": None,
                "required": False,
            },
        ],
    ),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
BLOCK_FILL = PatternFill("solid", fgColor="D6E4F0")
BLOCK_FONT = Font(bold=True, size=11, color="1F4E79")
THIN = Side(style="thin", color="B4B4B4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def style_header_row(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def build_questionnaire_sheet(wb):
    ws = wb.active
    ws.title = "Опросник"
    ws["A1"] = "Опросник для директоров и руководителей филиалов"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws.merge_cells("A1:E1")
    ws["A2"] = (
        "Проект: диагностика процессов ДУП (Департамент управления проектами), ТОО «Көркем Телеком». "
        "Цель: выявить боли и проблемы во взаимодействии филиал ↔ ДУП ↔ ЦО. "
        "Время заполнения: 25–35 минут. Анонимность: по решению организатора."
    )
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:E2")
    ws.row_dimensions[2].height = 45

    headers = ["№", "ID", "Блок", "Вопрос", "Тип ответа", "Варианты ответов", "Обязательный"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=4, column=i, value=h)
    style_header_row(ws, 4, len(headers))

    row = 5
    n = 1
    for block_id, block_title, questions in BLOCKS:
        ws.cell(row=row, column=1, value="")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        bc = ws.cell(row=row, column=1, value=block_title)
        bc.fill = BLOCK_FILL
        bc.font = BLOCK_FONT
        bc.alignment = WRAP
        bc.border = BORDER
        for c in range(2, 8):
            ws.cell(row=row, column=c).fill = BLOCK_FILL
            ws.cell(row=row, column=c).border = BORDER
        row += 1

        for q in questions:
            opts = q["options"]
            if q["type"] == "матрица":
                opt_text = "Для каждого подразделения: шкала 1–5 (1=плохо, 5=отлично). Подразделения: " + "; ".join(opts)
            elif opts:
                opt_text = " | ".join(opts)
            else:
                opt_text = "—"

            values = [
                n,
                q["id"],
                block_id,
                q["text"],
                q["type"],
                opt_text,
                "Да" if q["required"] else "Нет",
            ]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.alignment = WRAP
                cell.border = BORDER
            ws.row_dimensions[row].height = max(30, min(120, 15 * (1 + len(q["text"]) // 80)))
            row += 1
            n += 1

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 55
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 50
    ws.column_dimensions["G"].width = 12
    ws.freeze_panes = "A5"


def build_response_sheet(wb):
    ws = wb.create_sheet("Сбор ответов")
    flat_questions = []
    for _, _, questions in BLOCKS:
        for q in questions:
            if q["type"] == "матрица":
                for sub in q["options"]:
                    flat_questions.append((f"{q['id']}_{sub[:20]}", f"{q['text']} [{sub}]"))
            else:
                flat_questions.append((q["id"], q["text"]))

    meta = ["Дата", "Регион/филиал", "Должность", "ФИО (опционально)"]
    headers = meta + [f"{qid}: {text[:60]}" for qid, text in flat_questions]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True, size=9)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = BORDER
    ws.row_dimensions[1].height = 80
    ws.freeze_panes = "E2"
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22 if col <= 4 else 18


def build_fill_sheet(wb):
    """Printable / fillable form for offline use."""
    ws = wb.create_sheet("Анкета для заполнения")
    ws["A1"] = "Анкета — директор / руководитель филиала (заполнить от руки или в Excel)"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws.merge_cells("A1:D1")
    ws["A2"] = "Дата: _____________   Филиал: _____________   Должность: _____________"
    ws.merge_cells("A2:D2")
    row = 4
    for block_id, block_title, questions in BLOCKS:
        ws.cell(row=row, column=1, value=block_title).font = BLOCK_FONT
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 1
        for q in questions:
            ws.cell(row=row, column=1, value=f"{q['id']}.")
            ws.cell(row=row, column=2, value=q["text"]).alignment = WRAP
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            ws.row_dimensions[row].height = max(25, 15 * (1 + len(q["text"]) // 70))
            row += 1
            if q["type"] == "матрица":
                for sub in q["options"]:
                    ws.cell(row=row, column=2, value=f"  • {sub}:  1  2  3  4  5")
                    row += 1
            elif q["options"] and q["type"] != "шкала 1–5":
                for opt in q["options"]:
                    ws.cell(row=row, column=2, value=f"  ☐ {opt}")
                    row += 1
            elif q["type"] == "шкала 1–5":
                ws.cell(row=row, column=2, value="  Ответ:  1   2   3   4   5")
                row += 1
            else:
                ws.cell(row=row, column=2, value="  Ответ: " + "_" * 60)
                row += 1
            row += 1
        row += 1
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 80


def build_google_forms_sheet(wb):
    ws = wb.create_sheet("Google Forms — настройка")
    intro = [
        ["Параметр", "Значение"],
        ["Название формы", "ДУП — опрос директоров и руководителей филиалов"],
        ["Описание", "Диагностика процессов Департамента управления проектами. Цель — выявить боли во взаимодействии филиал ↔ ДУП ↔ центральный офис. Время: 25–35 мин."],
        ["Сбор email", "По желанию (для приглашения на интервью)"],
        ["Ограничение", "1 ответ"],
        ["Перемешивание вопросов", "Нет"],
    ]
    for r, row_data in enumerate(intro, 1):
        for c, val in enumerate(row_data, 1):
            ws.cell(row=r, column=c, value=val).border = BORDER

    headers = ["Порядок", "ID", "Раздел (заголовок)", "Тип в Google Forms", "Текст вопроса", "Варианты / подсказка", "Обязательный"]
    start = 8
    for i, h in enumerate(headers, 1):
        ws.cell(row=start, column=i, value=h)
    style_header_row(ws, start, len(headers))

    gforms_type_map = {
        "текст": "Краткий ответ",
        "развёрнутый текст": "Абзац",
        "один из списка": "Один из списка",
        "несколько из списка": "Несколько флажков",
        "да/нет": "Один из списка",
        "шкала 1–5": "Шкала (1–5)",
        "матрица": "Несколько вопросов-шкал (создать отдельно на каждое подразделение)",
    }

    row = start + 1
    order = 1
    for block_id, block_title, questions in BLOCKS:
        ws.cell(row=row, column=1, value=order)
        ws.cell(row=row, column=3, value=f"▌ {block_title}")
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
        for c in range(1, 8):
            ws.cell(row=row, column=c).fill = BLOCK_FILL
            ws.cell(row=row, column=c).border = BORDER
        row += 1
        order += 1

        for q in questions:
            if q["type"] == "матрица":
                for sub in q["options"]:
                    vals = [
                        order,
                        f"{q['id']}",
                        "",
                        "Шкала (1–5)",
                        f"{q['text']} — {sub}",
                        "1 = очень плохо, 5 = отлично",
                        "Да" if q["required"] else "Нет",
                    ]
                    for col, val in enumerate(vals, 1):
                        cell = ws.cell(row=row, column=col, value=val)
                        cell.alignment = WRAP
                        cell.border = BORDER
                    row += 1
                    order += 1
            else:
                opts = " | ".join(q["options"]) if q["options"] else ""
                if q["type"] == "шкала 1–5" and q["options"]:
                    opts = "Метки: " + " / ".join(q["options"][:2]) + " …"
                vals = [
                    order,
                    q["id"],
                    "",
                    gforms_type_map.get(q["type"], q["type"]),
                    q["text"],
                    opts,
                    "Да" if q["required"] else "Нет",
                ]
                for col, val in enumerate(vals, 1):
                    cell = ws.cell(row=row, column=col, value=val)
                    cell.alignment = WRAP
                    cell.border = BORDER
                row += 1
                order += 1

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 55
    ws.column_dimensions["F"].width = 45
    ws.column_dimensions["G"].width = 12


def main():
    wb = Workbook()
    build_questionnaire_sheet(wb)
    build_response_sheet(wb)
    build_fill_sheet(wb)
    build_google_forms_sheet(wb)
    wb.save(OUTPUT)
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    main()
