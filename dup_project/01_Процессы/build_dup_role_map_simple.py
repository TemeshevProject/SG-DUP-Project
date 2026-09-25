#!/usr/bin/env python3
"""Упрощённая карта ролей ДУП — одна вкладка, пять колонок.

Данные берутся из подробной карты (build_dup_role_map.py), чтобы два файла
не расходились: правка вносится один раз, оба файла пересобираются.

Колонки: уровень, наименование, бизнес-владелец, роль ДУП, что делает ДУП.
Роль сворачивается до четырёх значений по позиции ДУП относительно результата
объекта (Методика KT_METHOD, Глава 5 — владение, Глава 8 §1 — SIPOC):

    объект принадлежит ДУП                                   → Владелец
    ДУП выполняет внутри объекта шаг, создающий результат     → Участник
    ДУП стоит на выходе объекта и принимает результат         → Клиент
    владелец объекта не закреплён                             → Не определён

Клиент — не смягчённый участник, а отдельная позиция: по Методике внутренний
заказчик устанавливает требования к результату и принимает его, но шагов объекта
не выполняет и за его метрики не отвечает. Спорные случаи, где ДУП и передаёт
входы, и получает выход, разводятся по собственному продукту ДУП: есть
измеримый продукт внутри объекта (ФЭМ, реестр платежей, заявка на закупку) —
участник; вклад ограничен требованиями и приёмкой — клиент.

Значение «Не определён» сохранено в справочнике, но в текущей версии карты не
встречается: все объекты владения закреплены за должностями. Значение нужно,
чтобы возврат «ничьей» зоны был виден в фильтре, а не растворился в «участнике».
"""

from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from build_dup_role_map import НЗ, PROCESSES

OUTPUT = Path(__file__).resolve().parent / "ДУП_роли_в_процессах_кратко.xlsx"

ВЛАДЕЛЕЦ = "Владелец"
УЧАСТНИК = "Участник"
КЛИЕНТ = "Клиент"
НЕ_ОПРЕДЕЛЁН = "Не определён"

ROLE_ORDER = [ВЛАДЕЛЕЦ, УЧАСТНИК, КЛИЕНТ, НЕ_ОПРЕДЕЛЁН]
ROLE_FILLS = {
    ВЛАДЕЛЕЦ: PatternFill("solid", fgColor="C6E9D5"),
    УЧАСТНИК: PatternFill("solid", fgColor="CFE0F7"),
    КЛИЕНТ: PatternFill("solid", fgColor="FFE3B3"),
    НЕ_ОПРЕДЕЛЁН: PatternFill("solid", fgColor="F8CBC4"),
}

# Объекты, на выходе которых стоит ДУП: результат создаёт другое СП, ДУП
# формулирует требования, принимает результат и использует его в своих процессах.
CLIENT_CODES = {
    "1.1.1",   # пакет документов приходит от ДРБ
    "1.2",     # подтверждение готовности продукта — создаёт ДРБ
    "1.2.1",
    "1.5.3",   # утверждённый бюджет приходит из Финблока
    "1.7.2",   # закупочную процедуру проводят Закупщики
    "1.7.3",   # приёмка поставки
    "1.9.3",   # решение об оплате и проведение платежа
    "2.3.2",   # утверждённый прогноз
    "Б.2.3",   # подтверждение финансового закрытия
    "Б.5.3",   # зарегистрированный филиал принимает ДУП, регистрирует Юр блок
    "Б.6",     # передача проекта от ДРБ
    "Б.7",     # оборудование от СТ
    "Б.8",     # согласованный договор от Юр блока
    "Б.9",     # оформленный штат от HR
    "Б.10",    # смонтированная ИТ-инфраструктура
}

LEVEL_INDENT = {"Процесс": 0, "Подпроцесс": 1, "Этап": 2}
LEVEL_FILL = PatternFill("solid", fgColor="EDF1F7")  # подсветка строк уровня «Процесс»

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADERS = [
    "Уровень",
    "Наименование процесса / подпроцесса / этапа",
    "Бизнес-владелец процесса / подпроцесса / этапа",
    "Роль ДУП",
    "Что делает ДУП",
]
WIDTHS = [13, 54, 36, 15, 72]
NAME_COL = 2
ROLE_COL = 4


def simple_role(code: str, dup_role: str, owner_dept: str) -> str:
    """Сворачивает роль до четырёх значений по позиции ДУП относительно результата.

    Решающим является владелец процесса целиком, а не роль ДУП внутри своего
    дерева. Подтверждение готовности продукта (1.2) — подпроцесс блока 1, но
    результат создаёт ДРБ и принимает его ДУП, поэтому ДУП здесь клиент.
    """
    if dup_role == НЗ:
        return НЕ_ОПРЕДЕЛЁН
    if owner_dept.startswith("ДУП"):
        return ВЛАДЕЛЕЦ
    return КЛИЕНТ if code in CLIENT_CODES else УЧАСТНИК


def owner_text(dup_role: str, owner_dept: str, owner_role: str) -> str:
    """Бизнес-владелец одной строкой: подразделение и конкретная роль.

    Название подразделения подставляется, если роль сама по себе его не
    называет («Зам. по реализации» → «ДУП — Зам. по реализации»).
    """
    if dup_role == НЗ:
        return "не назначен"
    # При разделённом владении роль уже называет обе стороны — префикс лишний.
    if owner_dept == "ДУП" and "ДУП" not in owner_role:
        return f"ДУП — {owner_role}"
    return owner_role


def build_rows():
    rows = []
    for p in PROCESSES:
        code, level, name = p[0], p[1], p[2]
        owner_dept, owner_role, dup_role, contribution = p[6], p[7], p[8], p[10]
        rows.append((
            level,
            f"{code}  {name}" if code else name,
            owner_text(dup_role, owner_dept, owner_role),
            simple_role(code, dup_role, owner_dept),
            contribution.rstrip(" ."),
        ))
    return rows


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Роли ДУП"

    ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    rows = build_rows()
    for r, values in enumerate(rows, 2):
        level, role = values[0], values[ROLE_COL - 1]
        indent = LEVEL_INDENT.get(level, 0)
        for c, val in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = BORDER
            cell.alignment = Alignment(
                wrap_text=True, vertical="top", indent=indent if c == NAME_COL else 0
            )
        if level == "Процесс":
            ws.cell(row=r, column=1).fill = LEVEL_FILL
            ws.cell(row=r, column=NAME_COL).fill = LEVEL_FILL
            ws.cell(row=r, column=NAME_COL).font = Font(bold=True, size=10)
        role_cell = ws.cell(row=r, column=ROLE_COL)
        role_cell.fill = ROLE_FILLS[role]
        role_cell.font = Font(bold=True, size=10)

    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(rows) + 1}"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    wb.save(OUTPUT)

    counts = Counter(r[ROLE_COL - 1] for r in rows)
    print(f"Сохранено: {OUTPUT}")
    print(f"Записей: {len(rows)}")
    for role in ROLE_ORDER:
        print(f"  {role}: {counts.get(role, 0)}")


if __name__ == "__main__":
    main()
