"""Фирменный стиль Sergek Group (тема SG_pptx) для генерации презентаций.

Значения извлечены из корпоративного шаблона «Общие слайды_ver 12.12.24.pptx»:
слайд 16:9 (9144000 x 5143500 EMU), заголовок Golos Text SemiBold 20pt,
основной текст Golos Text 9–10pt, карточки на заливке F3F3F3.
"""

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

# --- Палитра темы SG_pptx ---
GREEN = RGBColor(0x28, 0xAB, 0x6A)  # accent1 — основной акцент
BLUE = RGBColor(0x01, 0x3D, 0x85)  # lt2 — корпоративный синий
BLUE_LIGHT = RGBColor(0x43, 0x84, 0xF4)  # accent5
YELLOW = RGBColor(0xFB, 0xBB, 0x05)  # accent4
RED = RGBColor(0xE9, 0x42, 0x35)  # accent6
DARK = RGBColor(0x1E, 0x1E, 0x1E)  # dk1
TEXT = RGBColor(0x28, 0x28, 0x28)
GRAY = RGBColor(0x44, 0x47, 0x46)
STEEL = RGBColor(0xA7, 0xB6, 0xBD)  # accent3
MIST = RGBColor(0xD4, 0xD8, 0xE4)  # accent2
CARD_BG = RGBColor(0xF3, 0xF3, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL_DARK = RGBColor(0x1C, 0x33, 0x41)  # dk2

# --- Шрифты ---
FONT_BOLD = "Golos Text SemiBold"
FONT_MED = "Golos Text Medium"
FONT = "Golos Text"

# --- Геометрия слайда (EMU) ---
SLIDE_W = 9144000
SLIDE_H = 5143500
MARGIN_L = 397125
TITLE_TOP = 327775
TITLE_H = 313500
CONTENT_TOP = 950000
CONTENT_BOTTOM = 4620000
CONTENT_W = SLIDE_W - 2 * MARGIN_L
CONTENT_H = CONTENT_BOTTOM - CONTENT_TOP

# --- Индексы макетов в шаблоне ---
LAYOUT_DARK_GRAY = 0
LAYOUT_TITLE = 2  # белый фон + логотип + футер + номер страницы
LAYOUT_WHITE = 3
LAYOUT_BLACK = 4


AMBER = RGBColor(0xA8, 0x73, 0x00)  # затемнённый YELLOW для текста на светлом фоне


def on(fill):
    """Контрастный цвет текста для заливки: жёлтый акцент требует тёмного текста."""
    return DARK if fill == YELLOW else WHITE


def ink(accent):
    """Читаемый вариант акцентного цвета для текста на светлом фоне."""
    return AMBER if accent == YELLOW else accent


def set_text(
    frame,
    lines,
    *,
    size=10,
    font=FONT,
    color=TEXT,
    bold=False,
    align=PP_ALIGN.LEFT,
    space_after=0,
    line_spacing=1.15,
):
    """Заполняет текстовый фрейм списком строк с единым стилем."""
    if isinstance(lines, str):
        lines = [lines]
    frame.word_wrap = True
    for i, line in enumerate(lines):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.alignment = align
        para.line_spacing = line_spacing
        para.space_after = Pt(space_after)
        run = para.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return frame


def textbox(
    slide,
    x,
    y,
    w,
    h,
    lines,
    *,
    size=10,
    font=FONT,
    color=TEXT,
    bold=False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    space_after=0,
    line_spacing=1.15,
):
    box = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = box.text_frame
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    set_text(
        tf,
        lines,
        size=size,
        font=font,
        color=color,
        bold=bold,
        align=align,
        space_after=space_after,
        line_spacing=line_spacing,
    )
    return box


def title(slide, text, subtitle=None):
    """Заголовок слайда в фирменной позиции шаблона."""
    textbox(
        slide,
        MARGIN_L,
        TITLE_TOP,
        6084300,
        TITLE_H,
        text,
        size=20,
        font=FONT_BOLD,
        color=DARK,
    )
    if subtitle:
        textbox(
            slide,
            MARGIN_L,
            TITLE_TOP + TITLE_H + 40000,
            CONTENT_W,
            220000,
            subtitle,
            size=10,
            color=GRAY,
        )
    return slide


def rect(slide, x, y, w, h, *, fill=CARD_BG, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06):
    sh = slide.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except (IndexError, ValueError):
            pass
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Emu(110000)
    tf.margin_top = tf.margin_bottom = Emu(90000)
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    return sh


def bullets(shape, lines, *, size=9, color=TEXT, font=FONT, prefix="", spacing=7, line_spacing=1.2):
    """Добавляет строки в текстовый фрейм фигуры с левым выравниванием."""
    tf = shape.text_frame
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if (i == 0 and not tf.paragraphs[0].runs) else tf.add_paragraph()
        para.alignment = PP_ALIGN.LEFT
        para.line_spacing = line_spacing
        para.space_before = Pt(0 if i == 0 else spacing)
        run = para.add_run()
        run.text = prefix + line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return shape


def card(slide, x, y, w, h, heading, body_lines, *, accent=GREEN, heading_size=11, body_size=9):
    """Карточка на светлой заливке с цветным заголовком."""
    sh = rect(slide, x, y, w, h)
    tf = sh.text_frame
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.line_spacing = 1.1
    r = p.add_run()
    r.text = heading
    r.font.name = FONT_BOLD
    r.font.size = Pt(heading_size)
    r.font.color.rgb = accent
    for line in body_lines:
        para = tf.add_paragraph()
        para.alignment = PP_ALIGN.LEFT
        para.line_spacing = 1.25
        para.space_before = Pt(5)
        run = para.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(body_size)
        run.font.color.rgb = TEXT
    return sh


def stat(slide, x, y, w, number, caption, *, accent=GREEN, num_size=30, cap_size=9, h_num=430000):
    """Крупная цифра с подписью."""
    textbox(slide, x, y, w, h_num, number, size=num_size, font=FONT_BOLD, color=accent, line_spacing=1.0)
    textbox(slide, x, y + h_num, w, 420000, caption, size=cap_size, color=GRAY, line_spacing=1.2)


def accent_bar(slide, x, y, w, h, text, *, fill=BLUE, color=WHITE, size=10, align=PP_ALIGN.CENTER):
    sh = rect(slide, x, y, w, h, fill=fill, radius=0.12)
    tf = sh.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = tf.margin_bottom = 0
    set_text(tf, text, size=size, font=FONT_BOLD, color=color, align=align, line_spacing=1.1)
    return sh


def divider(slide, x, y, w, *, color=MIST, thickness=1):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(thickness * 9525))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()
    line.shadow.inherit = False
    return line


def table_grid(slide, x, y, w, col_widths, rows, *, header=True, row_h=300000, size=9, zebra=True):
    """Лёгкая таблица из прямоугольников — полный контроль над стилем."""
    total = sum(col_widths)
    cols = [int(w * c / total) for c in col_widths]
    cur_y = y
    for r_i, row in enumerate(rows):
        cur_x = x
        is_head = header and r_i == 0
        for c_i, cell in enumerate(row):
            if is_head:
                fill, col, fnt, sz = BLUE, WHITE, FONT_BOLD, size
            else:
                fill = CARD_BG if (zebra and r_i % 2 == 0) else WHITE
                col, fnt, sz = TEXT, FONT, size
            box = rect(slide, cur_x, cur_y, cols[c_i], row_h, fill=fill, shape=MSO_SHAPE.RECTANGLE)
            tf = box.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = Emu(80000)
            tf.margin_top = tf.margin_bottom = Emu(20000)
            set_text(
                tf,
                str(cell),
                size=sz,
                font=fnt,
                color=col,
                align=PP_ALIGN.LEFT,
                line_spacing=1.05,
            )
            cur_x += cols[c_i]
        cur_y += row_h
    return cur_y
